from pathlib import Path
import pypdfium2 as pdfium
import docx
from docx import Document
import openpyxl


# ---------- PDF ----------
def pdf_text(pdf_path: str | Path) -> str:
    """Извлекает текст из PDF-файла.
        Args:
            pdf_path: Путь к PDF-файлу (строка или Path объект).
        Returns:
            Строка с текстом, извлеченным со всех страниц PDF
    """
    with pdfium.PdfDocument(pdf_path) as pdf:
        return "\n".join(
            pdf[idx].get_textpage().get_text_range()
            for idx in range(len(pdf))
        )


# ---------- DOCX ----------
def docx_text(docx_path: str | Path) -> str:
    """Извлекает текст из DOCX-файла.
        Args:
            docx_path: Путь к DOCX-файлу (строка или Path объект).

        Returns:
            Строка с текстом из всех непустых параграфов документа
    """
    return "\n".join(p.text for p in docx.Document(docx_path).paragraphs if p.text.strip())


# ---------- XLSX ----------
def xlsx_text(xlsx_path: str | Path) -> str:
    """Извлекает текст из XLSX-файла в виде markdown-таблиц (по одной на лист).
        Args:
            xlsx_path: Путь к XLSX-файлу (строка или Path объект).

        Returns:
            Строка с содержимым всех листов книги в виде markdown-таблиц
    """
    workbook = openpyxl.load_workbook(xlsx_path, data_only=True, read_only=True)
    sheets_text = []
    for sheet in workbook.worksheets:
        rows = [
            [("" if cell is None else str(cell)) for cell in row]
            for row in sheet.iter_rows(values_only=True)
            if any(cell is not None for cell in row)
        ]
        if not rows:
            continue
        lines = [f"### {sheet.title}", "| " + " | ".join(rows[0]) + " |"]
        lines.append("| " + " | ".join(["---"] * len(rows[0])) + " |")
        for row in rows[1:]:
            lines.append("| " + " | ".join(row) + " |")
        sheets_text.append("\n".join(lines))
    return "\n\n".join(sheets_text)


# ---------- auto-route ----------
def document_to_txt(file_path: str | Path) -> str:
    """Определяет тип документа и извлекает текст в зависимости от расширения.
        Args:
            file_path: Путь к файлу документа (строка или Path объект).

        Returns:
            Текст документа
    """
    try:
        file_path = Path(file_path)
        suffix = file_path.suffix.lower()
        if suffix == ".pdf":
            result_txt = pdf_text(file_path)
            if result_txt == "":
                return "Это скан pdf файла, напиши пользователю, что не получилось считать текст"
            return result_txt
        elif suffix == ".docx":
            return docx_text(file_path)
        elif suffix == ".xlsx":
            result_txt = xlsx_text(file_path)
            if result_txt == "":
                return "Файл пуст, напиши пользователю, что не получилось считать текст"
            return result_txt
        else:
            return " напиши пользователю, Only .pdf, .docx and .xlsx supported"
    except Exception as e:
        return f"Возникла ошибка {e}"


def text_to_docx(text: str, file_path: str | Path ):
    doc = Document()

    # Добавляем текст в документ
    doc.add_paragraph(text)

    # Сохраняем документ
    doc.save(str(file_path))




# ---------- use ----------
# text = document_to_txt("files/ggg.pdf")  # or .docx
# messages = [("human", f"Summarise this:\n\n{text}")]
# print(text)
